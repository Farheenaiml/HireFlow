"""Deterministic text work. No LLM calls live in this module - that is the point.

extract      : PDF / DOCX / TXT -> plain text
redact       : best-effort PII removal so screening is blind
verify_quotes: proves an LLM quote actually exists in the source
group        : transparent, code-owned scoring and grouping
"""
from __future__ import annotations

import hashlib
import io
import re
from typing import Any

import difflib

try:
    from rapidfuzz import fuzz
except Exception:  # pragma: no cover - rapidfuzz is optional, difflib is the fallback
    fuzz = None


def _ratio(a: str, b: str) -> float:
    """Similarity 0-100. rapidfuzz when installed, stdlib difflib otherwise."""
    if fuzz is not None:
        return float(fuzz.ratio(a, b))
    return difflib.SequenceMatcher(None, a, b).ratio() * 100.0

# --------------------------------------------------------------------------- extract


def extract_text(filename: str, data: bytes) -> dict:
    name = (filename or "").lower()
    warnings: list[str] = []
    text, method, pages = "", "plain", 1

    if name.endswith(".pdf"):
        method = "pdf"
        try:
            import fitz  # PyMuPDF

            doc = fitz.open(stream=data, filetype="pdf")
            pages = doc.page_count
            text = "\n".join(p.get_text("text") for p in doc)
            doc.close()
        except Exception as exc:
            warnings.append(f"PDF parse failed ({exc}). Try DOCX or TXT.")
        if len(text.strip()) < 40:
            warnings.append("No text layer found in this PDF. Scanned resumes need OCR, "
                            "which is out of scope for this prototype.")
    elif name.endswith(".docx"):
        method = "docx"
        try:
            import docx  # python-docx

            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        except Exception as exc:
            warnings.append(f"DOCX parse failed ({exc}).")
    else:
        for enc in ("utf-8", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except Exception:
                continue

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return {"text": text, "pages": pages, "method": method, "warnings": warnings}


# --------------------------------------------------------------------------- redact

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+|00)?\d{0,3}[\s-]?(?:\(?\d{2,5}\)?[\s.-]?){2,4}\d{2,5}")
URL_RE = re.compile(r"(?:https?://|www\.)\S+|(?:linkedin\.com|github\.com)/\S+", re.I)
INSTITUTION_RE = re.compile(
    r"\b([A-Z][\w&.'-]*(?:\s+[A-Z][\w&.'-]*){0,4}\s+"
    r"(?:University|College|Institute|Institution|School|Academy|Polytechnic))\b|"
    r"\b(University|Institute)\s+of\s+[A-Z][\w'-]*(?:\s+[A-Z][\w'-]*)?",
)
SENSITIVE_LINE_RE = re.compile(
    r"\b(date of birth|d\.o\.b|dob|age\s*[:\-]|marital status|nationality|religion|caste|"
    r"gender\s*[:\-]|sex\s*[:\-]|passport|aadhaar|visa status|photo)\b", re.I)
ADDRESS_RE = re.compile(
    r"\b\d{1,5}\s+[A-Z][\w.'-]*(?:\s+[A-Z][\w.'-]*)*\s+"
    r"(Street|St\.|Road|Rd\.|Avenue|Ave\.|Lane|Ln\.|Block|Sector|Nagar|Colony)\b", re.I)
PRONOUNS = {"he": "they", "she": "they", "him": "them", "her": "their",
            "his": "their", "hers": "theirs", "He": "They", "She": "They",
            "His": "Their", "Her": "Their"}
NAME_STOPWORDS = {
    "resume", "curriculum", "vitae", "cv", "profile", "summary", "objective",
    "experience", "education", "skills", "projects", "contact", "software",
    "engineer", "developer", "manager", "senior", "junior", "lead", "data",
    "product", "full", "stack", "backend", "frontend",
}


def _looks_like_name(line: str) -> bool:
    words = line.strip().split()
    if not 1 < len(words) <= 4:
        return False
    if any(w.lower().strip(",.|") in NAME_STOPWORDS for w in words):
        return False
    if any(ch.isdigit() or ch in "@/" for ch in line):
        return False
    return all(w[:1].isupper() for w in words if w[:1].isalpha())


def redact(text: str) -> dict:
    pii_map: dict[str, list[str]] = {}
    counts: dict[str, int] = {}

    def note(kind: str, value: str) -> None:
        pii_map.setdefault(kind, [])
        if value and value not in pii_map[kind]:
            pii_map[kind].append(value)
        counts[kind] = counts.get(kind, 0) + 1

    lines = text.split("\n")
    # 1. candidate name from the header block (first 6 non-empty lines)
    names: list[str] = []
    seen = 0
    for i, line in enumerate(lines):
        if not line.strip():
            continue
        seen += 1
        if seen > 6:
            break
        if _looks_like_name(line):
            names.append(line.strip())
            lines[i] = "[CANDIDATE]"
            note("name", line.strip())
    text = "\n".join(lines)

    for nm in names:
        for token in nm.split():
            if len(token) > 2:
                text = re.sub(rf"\b{re.escape(token)}\b", "[CANDIDATE]", text)

    def sub(pattern, repl, kind):
        nonlocal text

        def _r(m):
            note(kind, m.group(0))
            return repl

        text = pattern.sub(_r, text)

    sub(EMAIL_RE, "[EMAIL]", "email")
    sub(URL_RE, "[LINK]", "link")
    sub(PHONE_RE, "[PHONE]", "phone")
    sub(ADDRESS_RE, "[ADDRESS]", "address")
    sub(INSTITUTION_RE, "[INSTITUTION]", "institution")

    kept = []
    for line in text.split("\n"):
        if SENSITIVE_LINE_RE.search(line):
            counts["sensitive_line"] = counts.get("sensitive_line", 0) + 1
            pii_map.setdefault("sensitive_line", []).append(line.strip())
            continue
        kept.append(line)
    text = "\n".join(kept)

    text = re.sub(r"\b(" + "|".join(PRONOUNS) + r")\b",
                  lambda m: PRONOUNS[m.group(0)], text)
    def _agree(m: re.Match) -> str:
        subj, verb = m.group(1), m.group(2)
        irregular = {"has": "have", "was": "were", "is": "are", "does": "do"}
        if verb.lower() in irregular:
            return f"{subj} {irregular[verb.lower()]}"
        if verb.endswith("ies") and len(verb) > 4:
            return f"{subj} {verb[:-3]}y"
        if verb.endswith(("sses", "xes", "zes", "ches", "shes")):
            return f"{subj} {verb[:-2]}"
        if verb.endswith("s") and not verb.endswith(("ss", "us", "is")):
            return f"{subj} {verb[:-1]}"
        return m.group(0)

    text = re.sub(r"\b(They|they) ([a-z]{2,}s)\b", _agree, text)

    return {
        "redacted_text": text,
        "pii_map": pii_map,
        "redactions": [{"type": k, "count": v} for k, v in sorted(counts.items())],
    }


# --------------------------------------------------------------------------- verify

_WS = re.compile(r"\s+")


def _norm(s: str) -> str:
    s = (s or "").replace("\u2019", "'").replace("\u2018", "'")
    s = s.replace("\u201c", '"').replace("\u201d", '"')
    s = s.replace("\u2013", "-").replace("\u2014", "-").replace("\u2022", " ")
    return _WS.sub(" ", s).strip().lower()


def _best_alignment(qn: str, src: str) -> tuple[float, int | None, int | None]:
    """Best fuzzy location of quote `qn` inside `src` -> (score 0-100, start, end)."""
    if not src:
        return 0.0, None, None
    if fuzz is not None and hasattr(fuzz, "partial_ratio_alignment"):
        al = fuzz.partial_ratio_alignment(qn, src)
        if al is not None:
            return float(al.score), int(al.dest_start), int(al.dest_end)
    best, best_span = 0.0, (None, None)
    n = len(qn)
    step = max(3, n // 12)
    for width in sorted({max(8, n - 8), n, n + 8}):
        for i in range(0, max(1, len(src) - width + 1), step):
            sc = _ratio(qn, src[i:i + width])
            if sc > best:
                best, best_span = sc, (i, min(len(src), i + width))
    return best, best_span[0], best_span[1]


def verify_quotes(source_text: str, items: list[dict]) -> list[dict]:
    """Every AI claim must point at real text. This is what makes it provable."""
    src_norm = _norm(source_text)
    results = []
    for item in items:
        qid = item.get("id")
        quote = (item.get("quote") or "").strip()
        res = {"id": qid, "verified": False, "method": "none", "score": 0.0,
               "start": None, "end": None, "matched_text": ""}
        if not quote:
            results.append(res)
            continue
        if len(quote) > 300:
            res["method"] = "too_long"
            results.append(res)
            continue
        qn = _norm(quote)
        pos = src_norm.find(qn)
        if pos >= 0:
            res.update(verified=True, method="exact", score=100.0, start=pos,
                       end=pos + len(qn), matched_text=source_text and qn)
            results.append(res)
            continue
        if len(qn) > 12:
            score, start, end = _best_alignment(qn, src_norm)
            if score >= 92 and start is not None:
                res.update(verified=True, method="fuzzy", score=round(score, 1),
                           start=start, end=end, matched_text=src_norm[start:end])
        results.append(res)
    return results


# --------------------------------------------------------------------------- score

STATUS_VALUE = {"met": 1.0, "partial": 0.5, "unclear": 0.25, "missing": 0.0}
CONFIG = {
    "weights": {"must": 0.8, "nice": 0.2},
    "status_value": STATUS_VALUE,
    "strong_must_score": 0.75,
    "potential_must_score": 0.45,
}


def group_candidate(requirements: list[dict], evaluations: list[dict],
                    config: dict | None = None) -> dict:
    cfg = {**CONFIG, **(config or {})}
    by_req = {r["req_id"]: r for r in requirements}
    must_vals, nice_vals = [], []
    counts = {"met": 0, "partial": 0, "unclear": 0, "missing": 0}
    missing_must, unclear_must, needs_validation = [], [], 0

    for ev in evaluations:
        status = ev.get("status", "missing")
        counts[status] = counts.get(status, 0) + 1
        if ev.get("needs_validation"):
            needs_validation += 1
        req = by_req.get(ev.get("req_id"))
        if not req:
            continue
        val = cfg["status_value"].get(status, 0.0)
        if req.get("priority") == "must":
            must_vals.append(val)
            if status == "missing":
                missing_must.append(ev["req_id"])
            if status == "unclear":
                unclear_must.append(ev["req_id"])
        else:
            nice_vals.append(val)

    must_score = sum(must_vals) / len(must_vals) if must_vals else 0.0
    nice_score = sum(nice_vals) / len(nice_vals) if nice_vals else 0.0
    score = 100 * (cfg["weights"]["must"] * must_score + cfg["weights"]["nice"] * nice_score)

    if must_score >= cfg["strong_must_score"] and not missing_must:
        grp = "Strong"
    elif must_score >= cfg["potential_must_score"] or must_score >= cfg["strong_must_score"]:
        grp = "Potential"
    else:
        grp = "Weak"

    met_musts = sum(1 for e in evaluations
                    if by_req.get(e.get("req_id"), {}).get("priority") == "must"
                    and e.get("status") == "met")
    total_musts = len(must_vals)
    rationale = (
        f"Meets {met_musts} of {total_musts} must-haves; "
        f"{len(unclear_must)} unclear; {len(missing_must)} missing; "
        f"{needs_validation} item(s) need validation."
    )
    return {
        "score": round(score, 1),
        "must_score": round(must_score, 3),
        "nice_score": round(nice_score, 3),
        "group": grp,
        "counts": counts,
        "needs_validation_count": needs_validation,
        "missing_must": missing_must,
        "rationale": rationale,
    }


def input_hash(*parts: Any) -> str:
    h = hashlib.sha256()
    for p in parts:
        h.update(str(p).encode("utf-8", "ignore"))
    return h.hexdigest()[:16]
